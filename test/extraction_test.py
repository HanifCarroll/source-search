import importlib.util, tempfile, unittest, json, zipfile
from pathlib import Path
spec=importlib.util.spec_from_file_location('extract',Path(__file__).parent.parent/'scripts/extract-documents.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
def fixture_pdf(path,pages,width=612,height=792):
 objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>'];kids=[]
 for text in pages:
  page_id=len(objects)+1;kids.append(f'{page_id} 0 R')
  stream=('BT /F1 12 Tf 72 700 Td ('+text+') Tj ET').encode()
  objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] /Resources << /Font << /F1 3 0 R >> >> /Contents {page_id+1} 0 R >>'.encode())
  objects.append(f'<< /Length {len(stream)} >>\nstream\n'.encode()+stream+b'\nendstream')
 objects[1]=f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {len(kids)} >>'.encode()
 data=bytearray(b'%PDF-1.4\n');offsets=[0]
 for i,obj in enumerate(objects,1):offsets.append(len(data));data.extend(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
 xref=len(data);data.extend(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
 for offset in offsets[1:]:data.extend(f'{offset:010} 00000 n \n'.encode())
 data.extend(f'trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode());path.write_bytes(data)

class ExtractionTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.source=self.root/'sources';self.source.mkdir();self.cache=self.root/'cache';self.cache.mkdir(mode=0o700)
 def test_same_named_attachments_keep_distinct_stable_references(self):
  rawdir=self.cache/'mail-raw';rawdir.mkdir(mode=0o700)
  raw=rawdir/'mail-fastmail-fixture.json';m.save(raw,{'source':'mail-fastmail','id':'fixture','raw':{},'accountSignature':'synthetic'})
  m.private_dir(self.cache/'attachments');folder=m.private_dir(self.cache/'attachments'/'mail-fastmail-fixture')
  first=folder/'one.txt';first.write_text('Invented first attachment')
  second=folder/'two.txt';second.write_text('Invented second attachment')
  entries=[{'path':str(first),'name':'note.txt'},{'path':str(second),'name':'note.txt'},{'path':str(first),'name':'note.txt'}]
  m.save(folder/'manifest.json',{'parent':m.signature(raw),'entries':entries})
  config={'maintenanceEnabled':True,'enabledSources':['documents','mail-fastmail'],'documentRoots':[],'libraryRoot':str(self.cache),'mailAccounts':{'mail-fastmail':{'account':'fixture'}}}
  m.main(config);records=json.loads((self.cache/'documents.json').read_text())['records']
  self.assertEqual(len(records),2);self.assertEqual(len({r['ref']['id'] for r in records}),2)
  self.assertEqual({r['body'].strip() for r in records},{'Invented first attachment','Invented second attachment'})
  m.main(config);self.assertEqual(json.loads((self.cache/'documents.json').read_text())['records'],records)
 def test_cached_takeout_entries_retain_the_selected_source(self):
  archive=self.source/'invented.zip'
  with zipfile.ZipFile(archive,'w') as z:z.writestr('fixture.txt','Invented export')
  folder=m.private_dir(self.cache/'takeout-attachments');attachment=folder/'fixture.txt';attachment.write_text('Invented attachment')
  entry={'path':str(attachment),'name':'fixture.txt','signature':m.signature(archive)+[1,0,0,'synthetic'],'locator':str(archive)+'#attachment/fixture'}
  m.save(folder/(m.digest(str(archive))+'.json'),{'parent':m.signature(archive),'entries':[entry]})
  records,failures=m.takeout_attachments({'enabledSources':['takeout-archive'],'takeout':[str(archive)]},self.cache)
  self.assertEqual(failures,[]);self.assertEqual(records,[{**entry,'source':'takeout-archive'}])
 def test_pdf_pages_docx_headings_and_html_no_script(self):
  from docx import Document
  p=self.source/'invented.pdf';fixture_pdf(p,['Synthetic violet comet','Invented second page']);text,_=m.extract(p,self.cache);self.assertIn('## Page 2',text);self.assertIn('violet comet',text)
  p=self.source/'invented.docx';d=Document();d.add_heading('Fixture heading',1);d.add_paragraph('Invented prose');d.save(p);text,_=m.extract(p,self.cache);self.assertIn('# Fixture heading',text);self.assertIn('Invented prose',text)
  p=self.source/'invented.html';p.write_text('<h1>Fixture</h1><p>Visible prose</p><script>hiddenCode</script>');text,_=m.extract(p,self.cache);self.assertIn('Visible prose',text);self.assertNotIn('hiddenCode',text)
 def test_snapshot_incremental_edit_delete_and_missing_root(self):
  p=self.source/'fixture.txt';p.write_text('Synthetic text');config={'maintenanceEnabled':True,'enabledSources':['documents'],'documentRoots':[str(self.source)],'libraryRoot':str(self.cache)};(self.cache/'mail-raw').mkdir()
  m.main(config);first=json.loads((self.cache/'documents.json').read_text());self.assertEqual(len(first['records']),1)
  m.main(config);again=json.loads((self.cache/'documents.json').read_text());self.assertEqual(first['records'],again['records'])
  p.write_text('Changed synthetic text');m.main(config);changed=json.loads((self.cache/'documents.json').read_text());self.assertNotEqual(first['records'][0]['ref']['version'],changed['records'][0]['ref']['version'])
  p.unlink();m.main(config);self.assertEqual(json.loads((self.cache/'documents.json').read_text())['records'],[])
  self.source.rmdir();self.assertRaises(ValueError,m.main,config)
 def test_export_files_data_tables_and_links_are_excluded(self):
  z=self.root/'fixture.zip'
  with zipfile.ZipFile(z,'w') as f:f.writestr('../../outside.txt','synthetic upload');f.writestr('file-upload.dat',b'%PDF-1.4 synthetic')
  (self.source/'link.txt').symlink_to(z);(self.source/'table.csv').write_text('a,b');(self.source/'table.tsv').write_text('a\tb')
  self.assertEqual(m.inventory({'documentRoots':[str(self.source)],**{k:[str(z)] for k in ('chatgpt','linkedin','takeout','twitter')}}),([],[]))
  self.assertFalse((self.root.parent/'outside.txt').exists())
 def test_book_length_pdf_is_skipped_and_stays_skipped(self):
  p=self.source/'invented-book.pdf';fixture_pdf(p,[f'Invented page {n}' for n in range(m.BOOK_PAGES+1)]);config={'maintenanceEnabled':True,'enabledSources':['documents'],'documentRoots':[str(self.source)],'libraryRoot':str(self.cache)};(self.cache/'mail-raw').mkdir()
  for _ in range(2):
   m.main(config);result=json.loads((self.cache/'documents.json').read_text())
   self.assertEqual((result['records'],result['failures']),([],[]));self.assertIn('page budget',result['skipped'][0]['reason'])
 def test_takeout_mime_attachment_keeps_message_locator_and_original_zip(self):
  from email.message import EmailMessage
  import mailbox,hashlib
  message=EmailMessage();message['From']='fixture@example.test';message['Subject']='Invented archive';message.set_content('Synthetic message body')
  message.add_attachment(b'Synthetic archived attachment',maintype='text',subtype='plain',filename='fixture.txt')
  mbox=self.root/'fixture.mbox';box=mailbox.mbox(mbox);box.add(message);box.flush();box.close()
  z=self.root/'takeout.zip'
  with zipfile.ZipFile(z,'w') as archive:archive.write(mbox,'Takeout/Mail/fixture.mbox')
  before=hashlib.sha256(z.read_bytes()).hexdigest();items,errors=m.takeout_attachments({'enabledSources':['takeout-archive'],'takeout':[str(z)]},self.cache)
  self.assertEqual(errors,[]);self.assertEqual(len(items),1);self.assertIn('/message/0/attachment/',items[0]['locator']);self.assertEqual(Path(items[0]['path']).read_text(),'Synthetic archived attachment')
  again,_=m.takeout_attachments({'enabledSources':['takeout-archive'],'takeout':[str(z)]},self.cache);self.assertEqual(items,again);self.assertEqual(before,hashlib.sha256(z.read_bytes()).hexdigest())
 def test_disabled_maintenance_never_reads_sources(self):
  with self.assertRaisesRegex(ValueError,'disabled'):m.main({'documentRoots':['/does-not-exist'],'libraryRoot':str(self.cache)})
 def test_disabled_mail_source_does_not_download_old_attachment(self):
  raw=self.cache/'mail-raw';raw.mkdir();(raw/'fixture.json').write_text(json.dumps({'source':'mail-personal','id':'fixture','raw':{'message':{'payload':{'filename':'fixture.txt','body':{'attachmentId':'synthetic','size':3}}}}}))
  old=m.command;m.command=lambda *a,**k:self.fail('disabled provider invoked')
  try:self.assertEqual(m.mail_attachments({'enabledSources':['documents']},self.cache),([],[]))
  finally:m.command=old
 def test_isolated_worker_enforces_output_bound(self):
  p=self.source/'large.txt';p.write_text('a'*(4*1024*1024+1))
  with self.assertRaisesRegex(ValueError,'4 MB'):m.isolated_extract(p,self.cache,{})
 def test_private_directory_refuses_symlink(self):
  link=self.root/'linked';link.symlink_to(self.cache)
  with self.assertRaisesRegex(ValueError,'links'):m.private_dir(link)
 def test_pdf_ocr_uses_rendered_page_and_cleans_temporary_image(self):
  p=self.source/'blank.pdf';fixture_pdf(p,['']);old=m.command;calls=[]
  def ocr(args,*a,**k):
   from PIL import Image
   with Image.open(args[1]) as image:self.assertGreater(image.width,100)
   calls.append(args);return b'Synthetic recognized text'
  m.command=ocr
  try:
   text,warnings=m.extract(p,self.cache,{'tesseractBinary':'/synthetic/ocr'})
   self.assertIn('Synthetic recognized text',text);self.assertTrue(warnings);self.assertEqual(calls[0][0],'/synthetic/ocr');self.assertEqual(list(self.cache.glob('ocr-*')),[])
  finally:m.command=old
 def test_huge_pdf_page_fails_before_render(self):
  p=self.source/'huge.pdf';fixture_pdf(p,[''],100000,100000)
  with self.assertRaisesRegex(ValueError,'pixel budget'):m.extract(p,self.cache)
if __name__=='__main__':unittest.main()
